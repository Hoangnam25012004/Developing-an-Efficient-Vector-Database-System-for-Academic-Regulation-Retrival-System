"""Fixture builders shared by the tests (all files are generated, none shipped)."""

from __future__ import annotations

import importlib
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import docx  # noqa: E402
import fitz  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402

ARIAL = r"C:\Windows\Fonts\arial.ttf"


def parser():
    return importlib.import_module("src.pipeline.01_parse_chunk")


def make_pdf(path: Path, pages: list[str], password: str | None = None) -> Path:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        kw = {"fontfile": ARIAL, "fontname": "arial"} if Path(ARIAL).exists() else {}
        page.insert_textbox(fitz.Rect(50, 50, 550, 800), text, fontsize=11, **kw)
    if password:
        doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password)
    else:
        doc.save(path)
    doc.close()
    return path


def make_docx(path: Path, build) -> Path:
    """build(document) fills a fresh python-docx Document."""
    d = docx.Document()
    build(d)
    d.save(path)
    return path


def add_custom_list(document) -> str:
    """Register a two-level list: level 0 "1.", level 1 "a)". Returns its numId."""
    numbering = document.part.numbering_part.element
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90">'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2)"/></w:lvl>'
        '</w:abstractNum>'
    )
    num = parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>')
    # abstractNum elements must precede num elements.
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(num)
    return "90"


def numbered_paragraph(document, text: str, num_id: str, ilvl: int = 0):
    p = document.add_paragraph(text)
    ppr = p._p.get_or_add_pPr()
    ppr.append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'
    ))
    return p


def add_run_xml(paragraph, xml: str) -> None:
    paragraph._p.append(parse_xml(xml))


def rewrite_zip(src: Path, dst: Path, replace: dict[str, bytes] | None = None,
                add: dict[str, bytes] | None = None) -> Path:
    replace = replace or {}
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = replace.get(item.filename, zin.read(item.filename))
            zout.writestr(item, data)
        for name, data in (add or {}).items():
            zout.writestr(name, data)
    return dst
