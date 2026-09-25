import tempfile
import unittest
from pathlib import Path

from docx.oxml.ns import nsdecls

from tests._util import add_custom_list, add_run_xml, make_docx, numbered_paragraph, parser
from src.ingest.docx_extract import extract_docx

pc = parser()
HELPERS = dict(caption_ok=pc._usable_caption, header_candidate=pc._is_header_candidate,
               build_labels=pc._build_column_labels)


class DocxExtractTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def extract(self, build):
        path = make_docx(self.dir / "t.docx", build)
        return extract_docx(path, **HELPERS)

    def texts(self, ex):
        return [b.text for b in ex.blocks]

    def test_automatic_numbering_is_rebuilt(self):
        def build(d):
            d.add_paragraph("Điều 1. Phạm vi")
            d.add_paragraph("Quy định này áp dụng cho nhân viên.", style="List Number")
            d.add_paragraph("Nhân viên phải nộp báo cáo.", style="List Number")
        ex = self.extract(build)
        self.assertEqual(self.texts(ex)[1:], ["1. Quy định này áp dụng cho nhân viên.",
                                              "2. Nhân viên phải nộp báo cáo."])

    def test_multilevel_numbering_restarts_under_a_new_parent(self):
        def build(d):
            nid = add_custom_list(d)
            numbered_paragraph(d, "Mục lớn một", nid, 0)
            numbered_paragraph(d, "ý thứ nhất", nid, 1)
            numbered_paragraph(d, "ý thứ hai", nid, 1)
            numbered_paragraph(d, "Mục lớn hai", nid, 0)
            numbered_paragraph(d, "ý mới", nid, 1)
        ex = self.extract(build)
        self.assertEqual(self.texts(ex), ["1. Mục lớn một", "a) ý thứ nhất", "b) ý thứ hai",
                                          "2. Mục lớn hai", "a) ý mới"])

    def test_tracked_changes_resolve_as_accept_all(self):
        def build(d):
            p = d.add_paragraph("Giữ ")
            add_run_xml(p, f'<w:ins {nsdecls("w")} w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z">'
                           '<w:r><w:t xml:space="preserve">CHÈN </w:t></w:r></w:ins>')
            add_run_xml(p, f'<w:del {nsdecls("w")} w:id="2" w:author="a" w:date="2026-01-01T00:00:00Z">'
                           '<w:r><w:delText>XOÁ </w:delText></w:r></w:del>')
            p.add_run("cuối.")
        ex = self.extract(build)
        self.assertEqual(self.texts(ex), ["Giữ CHÈN cuối."])

    def test_headings_carry_their_outline_level(self):
        def build(d):
            d.add_heading("Giới thiệu", 1)
            d.add_paragraph("Nội dung.")
            d.add_heading("Phạm vi", 2)
        ex = self.extract(build)
        kinds = [(b.kind, b.level) for b in ex.blocks]
        self.assertEqual(kinds, [("heading", 0), ("para", None), ("heading", 1)])

    def test_table_rows_become_labelled_chunks(self):
        def build(d):
            d.add_paragraph("Bảng 1. Mức hỗ trợ công tác phí")
            t = d.add_table(rows=3, cols=3)
            for r, row in enumerate([["Hạng mục", "Mức", "Ghi chú"],
                                     ["Vé máy bay", "Theo thực tế", ""],
                                     ["Lưu trú", "800.000 đồng/đêm", "Tối đa 3 đêm"]]):
                for c, v in enumerate(row):
                    t.cell(r, c).text = v
        ex = self.extract(build)
        self.assertEqual(len(ex.table_chunks), 2)
        first = ex.table_chunks[0]["text"].split("\n")
        self.assertEqual(first[0], "Bảng 1. Mức hỗ trợ công tác phí")
        self.assertEqual(first[1], "[Bảng 1 · trang 1 · dòng 1]")
        self.assertIn("Hạng mục: Vé máy bay", first)
        # A generic table (no numeric first column) is not dropped, unlike PDF tables.
        self.assertIn("Ghi chú: Tối đa 3 đêm", ex.table_chunks[1]["text"])
        # Table text is not duplicated into the page text.
        self.assertNotIn("Vé máy bay", "\n".join(p["text"] for p in ex.pages))
        self.assertIn("Vé máy bay", ex.linear_text)

    def test_merged_cells_are_read_once(self):
        def build(d):
            nid = add_custom_list(d)
            t = d.add_table(rows=3, cols=2)
            t.cell(0, 0).text, t.cell(0, 1).text = "Nhóm", "Nội dung"
            merged = t.cell(1, 0).merge(t.cell(2, 0))
            merged.paragraphs[0].text = ""
            p = merged.paragraphs[0]
            p.add_run("Nhóm A")
            t.cell(1, 1).text, t.cell(2, 1).text = "x", "y"
            numbered_paragraph(d, "sau bảng", nid, 0)
        ex = self.extract(build)
        rows = [c["text"] for c in ex.table_chunks]
        self.assertEqual(len(rows), 2)
        self.assertTrue(all("Nhóm: Nhóm A" in r for r in rows))   # vertical merge repeats per row
        self.assertEqual(self.texts(ex)[-1], "1. sau bảng")      # numbering not advanced twice

    def test_layout_table_is_read_as_text(self):
        def build(d):
            t = d.add_table(rows=1, cols=2)
            t.cell(0, 0).text = "BỘ GIÁO DỤC VÀ ĐÀO TẠO\nSố: 12/QĐ-BGDĐT"
            t.cell(0, 1).text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
            d.add_paragraph("Điều 1. Nội dung")
        ex = self.extract(build)
        self.assertEqual(ex.table_chunks, [])
        self.assertIn("Số: 12/QĐ-BGDĐT", ex.pages[0]["text"])

    def test_page_numbers_follow_breaks(self):
        def build(d):
            d.add_paragraph("Trang một")
            d.add_page_break()
            d.add_paragraph("Trang hai")
        ex = self.extract(build)
        self.assertEqual([(b.text, b.page) for b in ex.blocks], [("Trang một", 1), ("Trang hai", 2)])

    def test_rendered_page_breaks_take_precedence(self):
        def build(d):
            d.add_paragraph("Một")
            d.add_page_break()                 # explicit break: ignored in rendered mode
            p = d.add_paragraph()
            add_run_xml(p, f'<w:r {nsdecls("w")}><w:lastRenderedPageBreak/><w:t>Hai</w:t></w:r>')
        ex = self.extract(build)
        self.assertEqual([(b.text, b.page) for b in ex.blocks], [("Một", 1), ("Hai", 2)])

    def test_skipped_parts_are_reported(self):
        def build(d):
            d.sections[0].header.paragraphs[0].text = "Tiêu đề trang"
            d.add_paragraph("Nội dung chính")
        ex = self.extract(build)
        self.assertIn("header_footer_skipped", ex.warnings)
        self.assertNotIn("Tiêu đề trang", ex.linear_text)


if __name__ == "__main__":
    unittest.main()
